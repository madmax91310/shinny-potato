"""Exact official total-return index histories for the anniversary format."""
import argparse
import base64
import csv
import datetime as dt
import hashlib
import io
import json
import pathlib
from concurrent.futures import ThreadPoolExecutor
from data_automation import UTC, month_shift, number, reject, write_json_atomic
from issuer_documents import download
from refresh_monthly_history import collect_one as collect_monthly

ROOT = pathlib.Path(__file__).resolve().parents[1]
HISTORY = ROOT/'src/data/anniversary-index-history.json'
PROOF = ROOT/'scripts/source-snapshots/anniversary-indices.json'
RUSSELL_COLUMNS = ['Index_Name','Date','Value_Without_Dividends__USD_','Value_With_Dividends__USD_']

def configurations():
    return json.loads((ROOT/'scripts/anniversary-indices.json').read_text())['series']

def russell_rows(raw):
    reader = csv.DictReader(io.StringIO(raw.decode('utf-8-sig')))
    if reader.fieldnames != RUSSELL_COLUMNS: reject('Russell currency/return columns changed')
    rows = []
    for row in reader:
        if row['Index_Name'] != 'Russell 2000®': reject('Wrong Russell index')
        day = dt.datetime.strptime(row['Date'], '%m/%d/%Y').date()
        rows.append((day, number(float(row[RUSSELL_COLUMNS[3]]))))
    dates = [day for day,_ in rows]
    if not dates or dates != sorted(set(dates), reverse=True): reject('Russell duplicate/unordered dates')
    return rows

def allworld_rows(raw):
    from openpyxl import load_workbook
    book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    try:
        if 'AWORLDS' not in book.sheetnames: reject('FTSE All-World sheet absent')
        sheet = book['AWORLDS']
        if sheet['B5'].value != 'FTSE All-World (USD)': reject('Wrong FTSE index/currency')
        if [sheet.cell(9,col).value for col in range(2,7)] != ['Date','Capital Return','Total Return','Capital Return','Total Return']:
            reject('FTSE price/total-return columns changed')
        if sheet['C8'].value != 'Index Value' or sheet['E8'].value != 'Index Return %': reject('FTSE levels/returns changed')
        rows=[]
        for cells in sheet.iter_rows(min_row=10,values_only=True):
            if cells[1] is None: continue
            if not isinstance(cells[1],dt.datetime): reject('Invalid FTSE observation date')
            rows.append((cells[1].date(),number(cells[3])))
    finally:
        book.close()
    dates=[day for day,_ in rows]
    if not dates or dates!=sorted(set(dates)): reject('FTSE duplicate/unordered dates')
    if len({day.strftime('%Y-%m') for day in dates})!=len(dates): reject('FTSE duplicate months')
    return rows

def ftse_rows(config, fetch=download):
    raw={url:fetch(url) for url in config['sourceUrls']}
    if config['parser']=='russell':
        combined={}
        for document in raw.values():
            for day,value in russell_rows(document):
                if day in combined and abs(combined[day]-value)>1e-6: reject('Russell files disagree')
                combined[day]=value
        rows=sorted(combined.items())
    elif config['parser']=='ftse': rows=allworld_rows(next(iter(raw.values())))
    else: reject('Unsupported FTSE parser')
    return rows,raw

def collect_ftse(config,now,fetch=download):
    rows,raw=ftse_rows(config,fetch)
    boundary=dt.datetime(now.year,now.month,1,tzinfo=UTC)
    final={}
    for day,value in rows:
        if config['periodStart']<=day.strftime('%Y-%m') and day<boundary.date(): final[day.strftime('%Y-%m')]=(day,value)
    if not final: reject('No completed FTSE months')
    keys=sorted(final)
    # All-World is a rolling official two-year workbook; never invent older levels.
    expected_start=config['periodStart'] if config['parser']=='russell' else keys[0]
    current=dt.datetime.strptime(expected_start,'%Y-%m').replace(tzinfo=UTC);expected=[]
    while current<boundary:
        expected.append(current.strftime('%Y-%m'));current=month_shift(current,1)
    if keys!=expected: reject('Incomplete/stale FTSE monthly history')
    proof=[]
    for key in keys:
        day,value=final[key]
        end=month_shift(dt.datetime.strptime(key,'%Y-%m').replace(tzinfo=UTC),1).date()
        if (end-day).days>7: reject('FTSE month-end observation too early')
        proof.append({'date':day.isoformat(),'value':value})
    encoded={url:base64.b64encode(document).decode() for url,document in raw.items()}
    return {**config,'points':[[key,final[key][1]] for key in keys], 'proofRows':proof,
            'rawResponse':encoded,'responseSha256':hashlib.sha256(json.dumps(encoded,sort_keys=True).encode()).hexdigest(),
            'checkedAt':now.date().isoformat(),'periodEnd':keys[-1]}

def merge_record(config,record,old):
    if not old: return record
    for key in ['currency','parser','variant','indexCode','method','periodStart','sourceUrls']:
        if key in config and old.get(key)!=config.get(key): reject('Index convention changed: '+key)
    previous=dict(old['points']);latest=dict(record['points'])
    if record['periodEnd']<old['periodEnd']:reject('Index month regression')
    overlap=previous.keys()&latest.keys()
    if not overlap:reject('Index histories have no overlap')
    # Official corrections may be tiny; a rebased or different return series must fail.
    for key in overlap:
        if abs(previous[key]-latest[key])>max(.02,previous[key]*1e-5):reject('Index historical level changed: '+key)
    points=sorted({**previous,**latest}.items())
    for a,b in zip(points,points[1:]):
        if month_shift(dt.datetime.strptime(a[0],'%Y-%m'),1).strftime('%Y-%m')!=b[0]:reject('Index monthly gap')
    return {**record,'points':[[key,value] for key,value in points]}

def refresh(current,proof,now,collect=None):
    end=(dt.date(now.year,now.month,1)-dt.timedelta(days=1)).strftime('%Y-%m')
    merged=current.copy();evidence=proof.copy();report={'checkedAt':now.isoformat(),'successes':[],'errors':[],'contractSkipped':[]}
    def one(config):
        try:
            old=current.get(config['id'])
            if old and old['periodEnd']==end and all(old.get(key)==value for key,value in config.items()):
                return config['id'],None,None
            record=collect(config,now) if collect else (collect_monthly(config,now) if config['parser']=='msci' else collect_ftse(config,now))
            return config['id'],merge_record(config,record,old),None
        except Exception as e:return config['id'],None,str(e)
    with ThreadPoolExecutor(max_workers=3) as pool:
        for id_,record,error in pool.map(one,configurations()):
            if error:report['errors'].append({'id':'anniversary-index-'+id_,'error':error});continue
            if record is None:
                report['contractSkipped'].append({'id':'anniversary-index-'+id_,'status':'cached','reason':'Completed month and source configuration already present; no network control in this invocation'})
                continue
            if record:
                merged[id_]={key:value for key,value in record.items() if key not in ('rawResponse','proofRows')}
                old_rows={r['date'][:7]:r for r in evidence.get(id_,{}).get('proofRows',[])}
                new_rows={r['date'][:7]:r for r in record['proofRows']}
                evidence[id_]={**record,'proofRows':[row for _,row in sorted({**old_rows,**new_rows}.items())]}
            report['successes'].append({'id':'anniversary-index-'+id_})
    return merged,evidence,report

def main():
    parser=argparse.ArgumentParser(description=__doc__);parser.add_argument('--apply',action='store_true');parser.add_argument('--output',required=True,type=pathlib.Path)
    args=parser.parse_args();old=json.loads(HISTORY.read_text()) if HISTORY.exists() else {};proof=json.loads(PROOF.read_text()) if PROOF.exists() else {}
    merged,evidence,report=refresh(old,proof,dt.datetime.now(UTC))
    write_json_atomic(args.output,report)
    if args.apply:
        if merged!=old:write_json_atomic(HISTORY,merged)
        if evidence!=proof:write_json_atomic(PROOF,evidence)
    print(json.dumps(report,ensure_ascii=False));return bool(report['errors'])

if __name__=='__main__':raise SystemExit(__import__('producer_report_contract').run('refresh_anniversary_indices', main))
