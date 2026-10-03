"""Update the shared gold series from World Bank's licensed monthly workbook."""
import argparse
import datetime as dt
import hashlib
import io
import json
import pathlib
import re
import time
import urllib.error
import urllib.parse
import urllib.request
import zipfile
from html.parser import HTMLParser

from openpyxl import load_workbook
from data_automation import number, write_json_atomic

PAGE = 'https://www.worldbank.org/en/research/commodity-markets'
CATALOG = 'https://datacatalog.worldbank.org/search/dataset/0038238/commodity-prices-history-and-projections'
LICENSE = 'https://creativecommons.org/licenses/by/4.0/'
TARGET = pathlib.Path(__file__).resolve().parents[1] / 'src/data/worldbank-gold-monthly.json'


def valid_url(url, workbook=False):
    parsed = urllib.parse.urlsplit(url)
    if parsed.scheme != 'https' or parsed.username or parsed.password or parsed.port:
        raise ValueError('Untrusted download URL')
    if workbook:
        if parsed.netloc != 'thedocs.worldbank.org' or not re.fullmatch(
                r'/en/doc/[A-Za-z0-9-]+/related/CMO-Historical-Data-Monthly\.xlsx', parsed.path):
            raise ValueError('Not the official monthly workbook')
    elif url != PAGE:
        raise ValueError('Not the official commodity page')


def download(url, workbook=False, opener=urllib.request.urlopen, sleep=time.sleep):
    valid_url(url, workbook)
    maximum = 5_000_000 if workbook else 2_000_000
    request = urllib.request.Request(url, headers={'User-Agent': 'EpargnantLibre-GoldRefresh/1.0'})
    for attempt in range(3):
        try:
            with opener(request, timeout=25) as response:
                valid_url(response.geturl(), workbook)
                body = response.read(maximum + 1)
                if len(body) > maximum:
                    raise ValueError('Download exceeds size limit')
                if workbook and not body.startswith(b'PK'):
                    raise ValueError('Expected XLSX, not an HTML error page')
                return body
        except urllib.error.HTTPError as error:
            if error.code not in (429, 500, 502, 503, 504) or attempt == 2:
                raise
        except (urllib.error.URLError, TimeoutError):
            if attempt == 2:
                raise
        sleep(2 ** attempt)


class Links(HTMLParser):
    def __init__(self):
        super().__init__()
        self.urls = set()

    def handle_starttag(self, tag, attrs):
        if tag == 'a':
            href = dict(attrs).get('href', '')
            if urllib.parse.urlsplit(href).path.endswith('/CMO-Historical-Data-Monthly.xlsx'):
                url = urllib.parse.urljoin(PAGE, href)
                valid_url(url, True)
                self.urls.add(url)


def discover_workbook(html):
    parser = Links()
    parser.feed(html)
    if len(parser.urls) != 1:
        raise ValueError('Expected one unambiguous monthly workbook link')
    return next(iter(parser.urls))


def month_range(end):
    year, month = map(int, end.split('-'))
    return [f'{i // 12:04d}-{i % 12 + 1:02d}' for i in range(2015 * 12, year * 12 + month)]


def read_gold(body, today):
    with zipfile.ZipFile(io.BytesIO(body)) as archive:
        if sum(entry.file_size for entry in archive.infolist()) > 30_000_000:
            raise ValueError('Expanded workbook exceeds size limit')
    workbook = load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    try:
        sheet = workbook['Monthly Prices']
        if sheet.max_row > 5000 or sheet.max_column > 200:
            raise ValueError('Unexpected workbook dimensions')
        rows = list(sheet.values)
        if rows[0][0] != 'World Bank Commodity Price Data (The Pink Sheet)':
            raise ValueError('Unexpected workbook identity')
        updated = re.fullmatch(r'Updated on ([A-Za-z]+ \d{2}, \d{4})', str(rows[3][0]))
        if not updated:
            raise ValueError('Missing publication date')
        published = dt.datetime.strptime(updated[1], '%B %d, %Y').date()
        if published > today:
            raise ValueError('Future publication date')
        columns = [i for i, cell in enumerate(rows[4]) if cell == 'Gold']
        if len(columns) != 1 or rows[5][columns[0]] != '($/troy oz)':
            raise ValueError('Gold column or unit changed')
        description_rows = [row for row in workbook['Description'].values
                            if len(row) > 1 and str(row[1]).startswith('Gold,')]
        if len(description_rows) != 1:
            raise ValueError('Missing gold methodology')
        description = description_rows[0][1]
        providers = str(description_rows[0][19])
        points = []
        for row in rows[6:]:
            if row[0] is None:
                continue
            match = re.fullmatch(r'(\d{4})M(0[1-9]|1[0-2])', str(row[0]))
            if not match:
                raise ValueError('Unexpected monthly date')
            month = f'{match[1]}-{match[2]}'
            if month < '2015-01':
                continue
            if month >= today.strftime('%Y-%m'):
                raise ValueError('Incomplete or future month in workbook')
            points.append([month, number(row[columns[0]])])
        if not points or [p[0] for p in points] != month_range(points[-1][0]):
            raise ValueError('Missing, duplicate or unordered gold months')
        if points[-1][0] >= published.strftime('%Y-%m'):
            raise ValueError('Price month inconsistent with publication date')
        return points, published.isoformat(), description, providers
    finally:
        workbook.close()


def prepare(previous, body, url, today):
    valid_url(url, True)
    points, published, description, providers = read_gold(body, today)
    if description != previous['seriesDescription'] or providers != previous['dataProviders']:
        raise ValueError('Source methodology changed; review required before updating')
    old = dict(previous['points'])
    new = dict(points)
    if any(month not in new for month in old):
        raise ValueError('Source truncated the active history')
    if any(abs(new[month] / value - 1) > .1 for month, value in old.items()):
        raise ValueError('Historical revision exceeds 10%; review required')
    if published < previous['workbookUpdatedAt']:
        raise ValueError('Publication regressed')
    result = {**previous, 'url': url, 'checkedAt': today.isoformat(),
              'workbookUpdatedAt': published, 'workbookSha256': hashlib.sha256(body).hexdigest(),
              'points': points}
    # A previous manual PDF cross-check must never certify a newer workbook.
    result.pop('crossCheck', None)
    changed = (points != previous['points'] or result['workbookSha256'] != previous['workbookSha256']
               or url != previous['url'])
    return result, changed


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', type=pathlib.Path, default=TARGET)
    parser.add_argument('--workbook', type=pathlib.Path, help='Replay a downloaded XLSX offline')
    parser.add_argument('--url', help='Official workbook URL for offline replay')
    args = parser.parse_args()
    today = dt.datetime.now(dt.timezone.utc).date()
    previous = json.loads(args.output.read_text())
    url = args.url if args.workbook else discover_workbook(download(PAGE).decode('utf-8'))
    if not url:
        parser.error('--workbook requires --url')
    body = args.workbook.read_bytes() if args.workbook else download(url, True)
    result, changed = prepare(previous, body, url, today)
    if changed:
        write_json_atomic(args.output, result)
    status = 'updated' if changed else 'unchanged'
    print(f'Gold {status}: {len(result["points"])} months, through {result["points"][-1][0]}')
    if result['points'][-1][0] != (today.replace(day=1) - dt.timedelta(days=1)).strftime('%Y-%m'):
        print('Latest complete month not published yet; existing valid history retained.')


if __name__ == '__main__':
    main()
